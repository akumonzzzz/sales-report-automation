"""Tests for sales_report.py. Run with: pytest"""

import json
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import sales_report as sr  # noqa: E402

CONFIG = sr.load_config(None)


def make_raw(rows):
    """Build the table load_files() would return from (order_id, date, product, qty, price, region) rows."""
    frame = pd.DataFrame(rows, columns=["order_id", "date", "product", "quantity", "unit_price", "region"])
    frame = frame.astype(object).where(frame.notna(), None)
    frame["source_file"] = "test.csv"
    frame["source_row"] = range(2, len(frame) + 2)
    return frame


# ---------------------------------------------------------------- parsing ---

@pytest.mark.parametrize("text", [
    "2026-03-12", "12/03/2026", "12-03-2026", "12/03/26", "12 Mar 2026", "12 March 2026",
    "Mar 12, 2026", "March 12, 2026", "2026/03/12", "2026-03-12 00:00:00", "  2026-03-12  ",
])
def test_parse_dates_accepts_every_supported_format(text):
    parsed = sr.parse_dates(pd.Series([text]), CONFIG["date_formats"])
    assert parsed.iloc[0] == pd.Timestamp(2026, 3, 12)


def test_parse_dates_is_day_first_by_default():
    parsed = sr.parse_dates(pd.Series(["03/04/2026"]), CONFIG["date_formats"])
    assert parsed.iloc[0] == pd.Timestamp(2026, 4, 3)


@pytest.mark.parametrize("text", ["not a date", "Feb 30, 2026", "31/02/2026", "", None])
def test_parse_dates_rejects_invalid_values(text):
    assert pd.isna(sr.parse_dates(pd.Series([text], dtype=object), CONFIG["date_formats"]).iloc[0])


@pytest.mark.parametrize("text, expected", [
    ("12", 12.0), (" 12 ", 12.0), ("29.90", 29.9), ("$29.90", 29.9), ("$1,299.50", 1299.5),
    ("AUD 45", 45.0), ("1,000,000", 1000000.0), ("-2", -2.0), ("(45.00)", -45.0), ("3 pcs", 3.0),
])
def test_parse_number_reads_common_formats(text, expected):
    assert sr.parse_number(pd.Series([text])).iloc[0] == pytest.approx(expected)


@pytest.mark.parametrize("text", ["", None, "abc", "1.299,50", "12,5", "1.2.3", "--5"])
def test_parse_number_rejects_anything_ambiguous(text):
    assert pd.isna(sr.parse_number(pd.Series([text], dtype=object)).iloc[0])


def test_parse_number_with_comma_decimal_separator():
    numbers = sr.parse_number(pd.Series(["1.299,50", "12,5", "€ 7"]), decimal_separator=",")
    assert list(numbers) == pytest.approx([1299.5, 12.5, 7.0])


def test_tidy_product_names_merges_case_variants_and_keeps_a_readable_spelling():
    names = pd.Series(["USB-C HUB", "usb-c hub", "USB-C Hub", "  usb-c   hub ", "MONITOR", "monitor", None, ""])
    tidy = sr.tidy_product_names(names)
    assert list(tidy[:4]) == ["USB-C Hub"] * 4
    assert list(tidy[4:6]) == ["Monitor", "Monitor"]  # no mixed-case spelling exists, so Title Case
    assert tidy[6:].isna().all()


# ---------------------------------------------------------------- loading ---

def test_normalise_columns_matches_aliases_ignoring_case_and_separators():
    frame = pd.DataFrame(columns=["Invoice-No", "SOLD_ON", "Product  Name", "units sold", "Price_Each", "State", "x"])
    out = sr.normalise_columns(frame, "file.csv", CONFIG["column_aliases"])
    assert list(out.columns) == ["order_id", "date", "product", "quantity", "unit_price", "region"]


def test_normalise_columns_adds_an_empty_order_id_when_the_file_has_none():
    frame = pd.DataFrame({"Date": ["2026-01-01"], "Product": ["A"], "Qty": ["1"], "Price": ["2"], "Region": ["NSW"]})
    assert sr.normalise_columns(frame, "file.csv", CONFIG["column_aliases"])["order_id"].isna().all()


def test_normalise_columns_names_the_missing_column():
    frame = pd.DataFrame(columns=["Date", "Product", "Qty", "Region"])
    with pytest.raises(ValueError, match="unit_price"):
        sr.normalise_columns(frame, "file.csv", CONFIG["column_aliases"])


@pytest.mark.parametrize("delimiter", [",", ";", "\t", "|"])
def test_read_source_detects_the_delimiter(tmp_path, delimiter):
    path = tmp_path / "sales.csv"
    path.write_text(delimiter.join(["Date", "Product", "Qty", "Price", "Region"]) + "\n"
                    + delimiter.join(["2026-01-01", "Mouse", "2", "29.90", "NSW"]) + "\n", encoding="utf-8")
    frame = sr.read_source(path)
    assert list(frame.columns) == ["Date", "Product", "Qty", "Price", "Region"]
    assert frame.iloc[0]["Price"] == "29.90"


def test_read_source_keeps_text_such_as_NA_and_reads_windows_encoding(tmp_path):
    path = tmp_path / "sales.csv"
    path.write_bytes("Date,Product,Qty,Price,Region\n2026-01-01,Café Mug,1,9.50,NA\n".encode("cp1252"))
    frame = sr.read_source(path)
    assert frame.iloc[0]["Region"] == "NA"       # North America, not a missing value
    assert frame.iloc[0]["Product"] == "Café Mug"


def test_load_files_skips_unreadable_files_and_numbers_rows_as_in_the_file(tmp_path):
    (tmp_path / "good.csv").write_text(
        "Date,Product,Qty,Price,Region\n2026-01-01,Mouse,2,29.90,NSW\n\n2026-01-02,Mouse,1,29.90,NSW\n",
        encoding="utf-8",
    )
    (tmp_path / "broken.csv").write_text("a,b\n1,2\n", encoding="utf-8")
    raw, skipped = sr.load_files(sr.find_input_files(tmp_path), CONFIG)
    assert list(raw["source_row"]) == [2, 4]     # the blank line 3 is skipped but still counted
    assert [name for name, _ in skipped] == ["broken.csv"]


def test_find_input_files_ignores_temp_files_and_the_report_itself(tmp_path):
    for name in ["a.csv", "b.xlsx", "~$b.xlsx", ".hidden.csv", "notes.txt", "report.xlsx"]:
        (tmp_path / name).write_text("x", encoding="utf-8")
    files = sr.find_input_files(tmp_path, output=tmp_path / "report.xlsx")
    assert [path.name for path in files] == ["a.csv", "b.xlsx"]


def test_find_input_files_fails_clearly_when_there_is_nothing_to_read(tmp_path):
    with pytest.raises(FileNotFoundError, match="No CSV or Excel files"):
        sr.find_input_files(tmp_path)


# --------------------------------------------------------------- cleaning ---

def test_clean_computes_revenue_and_normalises_values():
    good, rejected, duplicates = sr.clean(make_raw([[None, "12/03/2026", "usb-c hub", "2", "$59.50", " vic "]]))
    row = good.iloc[0]
    assert (row["date"], row["quantity"], row["unit_price"], row["revenue"], row["region"]) == (
        pd.Timestamp(2026, 3, 12), 2, 59.5, 119.0, "VIC")
    assert rejected.empty and duplicates.empty


@pytest.mark.parametrize("row, reason", [
    ([None, "", "Mouse", "1", "5", "NSW"], "missing date"),
    ([None, "Feb 30, 2026", "Mouse", "1", "5", "NSW"], "unrecognised date"),
    ([None, "2026-01-01", "", "1", "5", "NSW"], "missing product"),
    ([None, "2026-01-01", "Mouse", "", "5", "NSW"], "missing quantity"),
    ([None, "2026-01-01", "Mouse", "two", "5", "NSW"], "quantity is not a number"),
    ([None, "2026-01-01", "Mouse", "1.5", "5", "NSW"], "quantity is not a whole number"),
    ([None, "2026-01-01", "Mouse", "-2", "5", "NSW"], "quantity is zero or negative"),
    ([None, "2026-01-01", "Mouse", "0", "5", "NSW"], "quantity is zero or negative"),
    ([None, "2026-01-01", "Mouse", "1", "", "NSW"], "missing price"),
    ([None, "2026-01-01", "Mouse", "1", "1.299,50", "NSW"], "price is not a number"),
    ([None, "2026-01-01", "Mouse", "1", "-5", "NSW"], "negative price"),
    ([None, "", "TOTAL", "112", "", ""], "missing date; missing price"),
])
def test_clean_rejects_invalid_rows_with_the_reason(row, reason):
    good, rejected, _ = sr.clean(make_raw([row]))
    assert good.empty
    assert rejected.iloc[0]["reason"] == reason
    assert rejected.iloc[0]["source_row"] == 2


def test_clean_keeps_a_sale_with_no_region_and_labels_it():
    good, rejected, _ = sr.clean(make_raw([[None, "2026-01-01", "Mouse", "1", "5", ""]]))
    assert good.iloc[0]["region"] == sr.NO_REGION and rejected.empty


def test_clean_removes_identical_rows_and_says_which_row_they_repeat():
    sale = [None, "2026-01-01", "Mouse", "1", "5", "NSW"]
    good, _, duplicates = sr.clean(make_raw([sale, ["", "01/01/2026", "MOUSE", "1", "$5.00", "nsw"], sale]))
    assert len(good) == 1 and len(duplicates) == 2
    assert list(duplicates["duplicate_of"]) == ["test.csv row 2", "test.csv row 2"]


def test_clean_uses_order_ids_to_tell_repeat_sales_from_duplicates():
    rows = [
        ["INV-1", "2026-01-01", "Mouse", "1", "5", "NSW"],
        ["INV-1", "2026-01-01", "Mouse", "1", "5", "NSW"],   # same invoice line exported twice
        ["INV-2", "2026-01-01", "Mouse", "1", "5", "NSW"],   # a second, genuine sale
        ["INV-2", "2026-01-01", "Keyboard", "1", "9", "NSW"],  # another line on the same invoice
    ]
    good, _, duplicates = sr.clean(make_raw(rows))
    assert sorted(good["order_id"]) == ["INV-1", "INV-2", "INV-2"]
    assert list(duplicates["order_id"]) == ["INV-1"]


def test_clean_can_keep_duplicates():
    sale = [None, "2026-01-01", "Mouse", "1", "5", "NSW"]
    good, _, duplicates = sr.clean(make_raw([sale, sale]), keep_duplicates=True)
    assert len(good) == 2 and duplicates.empty


def test_every_row_is_accounted_for():
    rows = [
        [None, "2026-01-01", "Mouse", "1", "5", "NSW"],
        [None, "2026-01-01", "Mouse", "1", "5", "NSW"],
        [None, "bad", "Mouse", "1", "5", "NSW"],
        [None, "2026-02-01", "Keyboard", "3", "10", "VIC"],
    ]
    raw = make_raw(rows)
    good, rejected, duplicates = sr.clean(raw)
    assert len(good) + len(rejected) + len(duplicates) == len(raw)
    quality = sr.build_quality_table(raw, good, rejected, duplicates)
    assert quality.iloc[-1][["rows_read", "rows_kept", "rows_rejected", "duplicates_removed"]].tolist() == [4, 2, 1, 1]


# ------------------------------------------------------------ aggregation ---

def test_build_tables_totals_agree_with_each_other():
    rows = [
        [None, "2026-01-10", "Mouse", "2", "10", "NSW"],     # 20
        [None, "2026-01-20", "Keyboard", "1", "100", "VIC"],  # 100
        [None, "2026-02-05", "Mouse", "3", "10", "NSW"],     # 30
    ]
    good, _, _ = sr.clean(make_raw(rows))
    tables = sr.build_tables(good)

    assert tables["By Month"]["revenue"].tolist() == [120.0, 30.0]
    assert tables["By Month"]["growth_vs_prev_month"].iloc[1] == pytest.approx(-0.75)
    assert tables["By Product"]["product"].tolist() == ["Keyboard", "Mouse"]  # sorted by revenue
    assert tables["By Product"]["share_of_revenue"].sum() == pytest.approx(1.0)
    assert tables["By Region"].set_index("region")["revenue"].to_dict() == {"NSW": 50.0, "VIC": 100.0}
    assert tables["Region by Month"]["Total"].sum() == pytest.approx(150.0)
    for name in ["By Month", "By Product", "By Region"]:
        assert tables[name]["revenue"].sum() == pytest.approx(good["revenue"].sum())


# ----------------------------------------------------------------- config ---

def test_load_config_adds_aliases_and_replaces_date_formats(tmp_path):
    path = tmp_path / "config.json"
    path.write_text(json.dumps({
        "column_aliases": {"date": ["Datum"]}, "date_formats": ["%m/%d/%Y"], "currency_symbol": "€",
    }), encoding="utf-8")
    config = sr.load_config(path)
    assert config["column_aliases"]["date"][0] == "Datum"
    assert "order date" in config["column_aliases"]["date"]          # built-in aliases are kept
    assert config["date_formats"] == ["%m/%d/%Y"]
    assert sr.parse_dates(pd.Series(["03/04/2026"]), config["date_formats"]).iloc[0] == pd.Timestamp(2026, 3, 4)


@pytest.mark.parametrize("content, message", [
    ('{"colour": "blue"}', "unknown setting"),
    ('{"column_aliases": {"customer": ["Client"]}}', "unknown column"),
    ('{"decimal_separator": ";"}', "decimal_separator"),
    ('{"date_formats": []}', "date_formats"),
    ("{not json", "not valid JSON"),
])
def test_load_config_explains_mistakes(tmp_path, content, message):
    path = tmp_path / "config.json"
    path.write_text(content, encoding="utf-8")
    with pytest.raises(ValueError, match=message):
        sr.load_config(path)


# ------------------------------------------------------------- end to end ---

@pytest.fixture
def sample_folder(tmp_path):
    data = tmp_path / "data"
    data.mkdir()
    (data / "store.csv").write_text(
        "Date,Product,Qty,Unit Price,Region\n"
        "2026-01-05,Wireless Mouse,2,29.90,NSW\n"
        "2026-02-10,USB-C Hub,1,59.50,NSW\n", encoding="utf-8")
    (data / "web.csv").write_text(
        "Sold_On;Product_Name;Units;Price_Each;Location\n"
        "Jan 07, 2026;wireless mouse;1;29.90;Online\n"
        "Jan 07, 2026;wireless mouse;1;29.90;Online\n"
        "Feb 30, 2026;usb-c hub;1;59.50;Online\n", encoding="utf-8")
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Invoice No", "Invoice Date", "Description", "Units Sold", "Unit Price", "State"])
    sheet.append(["INV-1", pd.Timestamp(2026, 3, 1).to_pydatetime(), "USB-C Hub", 4, 59.5, "QLD"])
    sheet.append([None, None, "TOTAL", 4, None, None])
    workbook.save(data / "accounts.xlsx")
    return data


def test_full_run_writes_a_consistent_report(sample_folder, tmp_path):
    output = tmp_path / "out" / "report.xlsx"
    result = sr.run(sample_folder, output)

    assert (result.files_read, result.rows_read) == (3, 7)
    assert (result.rows_kept, result.rows_rejected, result.duplicates_removed) == (4, 2, 1)
    assert result.total_revenue == pytest.approx(2 * 29.90 + 59.50 + 29.90 + 4 * 59.50)

    workbook = load_workbook(output)
    assert workbook.sheetnames == [
        "Summary", "By Month", "By Product", "By Region", "Region by Month",
        "Clean Data", "Data Quality", "Rejected Rows", "Duplicates Removed",
    ]
    summary = {row[0]: row[1] for row in workbook["Summary"].iter_rows(max_col=2, values_only=True) if row[0]}
    assert summary["Total revenue"] == pytest.approx(result.total_revenue)
    assert summary["Transactions"] == 4
    assert summary["Top product by revenue"] == "USB-C Hub"
    assert len(workbook["Summary"]._charts) == 4

    clean_rows = list(workbook["Clean Data"].iter_rows(values_only=True))
    assert clean_rows[0][:3] == ("Order ID", "Date", "Product")
    assert sum(row[5] for row in clean_rows[1:]) == pytest.approx(result.total_revenue)
    reasons = [row[2] for row in workbook["Rejected Rows"].iter_rows(min_row=2, values_only=True)]
    assert sorted(reasons) == ["missing date; missing price", "unrecognised date"]


def test_report_is_the_same_on_every_run(sample_folder, tmp_path):
    def values(path):
        workbook = load_workbook(path)
        return {ws.title: [row for row in ws.iter_rows(values_only=True)] for ws in workbook.worksheets}

    sr.run(sample_folder, tmp_path / "first.xlsx")
    sr.run(sample_folder, tmp_path / "second.xlsx")
    assert values(tmp_path / "first.xlsx") == values(tmp_path / "second.xlsx")


def test_command_line_returns_0_on_success_and_1_on_error(sample_folder, tmp_path):
    assert sr.main(["--input", str(sample_folder), "--output", str(tmp_path / "r.xlsx"), "--quiet"]) == 0
    assert (tmp_path / "r.xlsx").exists()
    assert sr.main(["--input", str(tmp_path / "missing"), "--output", str(tmp_path / "x.xlsx"), "--quiet"]) == 1
    assert not (tmp_path / "x.xlsx").exists()
